"""Builds demo_workbook.xlsx - a small FABRICATED PinMux workbook (project DEMO-style layout, labelled CCS20 because the column
labels are fixed by the schema). No customer data. Every pin below is invented, so you know the right answer in advance:
see GUIDE.md for the expected result of each pin."""
import os, sys
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "backend"))
import openpyxl
from openpyxl.utils import get_column_letter
from app.schema import KNOWN_COLUMNS

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Database"
labels = list(KNOWN_COLUMNS.keys()); col = {}
for i, lab in enumerate(labels, 1):
    d = KNOWN_COLUMNS[lab]
    ws.cell(row=2, column=i, value=d.entity if d.entity else ("Parameter Class" if lab == "_PARA_ROW" else "_END_SB"))
    ws.cell(row=3, column=i, value=lab); col[lab] = i

def row(r, n, iface, group, sig, direction, pull, drv, func, remark=None):
    vals = {"SA8155SIP Ball Name (SA8155P)": f"GPIO_{n}", "SA8155SIP Ball ID (SA8155P)": f"X{n:03d}", "Power Domain (SA8155P)": "VDD_1P8",
            "RegSetting1 (PinMux01)": f"TLMM_GPIO_CFG{n}(5:2)={func}",
            "PinMuxed Name (CCS20)": f"GPIO{n}", "Selected Interface (CCS20)": iface, "Interface Group (CCS20)": group,
            "Gen5 Signal Name (CCS20)": sig, "Direction (CCS20)": direction, "IO_GPIO_PULL (CCS20)": pull,
            "IO_DRV_STRENGTH (CCS20)": drv, "ResetValue (CCS20)": "0: low"}
    if remark: vals["Remarks (CCS20)"] = remark
    for k, v in vals.items(): ws.cell(row=r, column=col[k], value=v)

R = 4
PINS = [
 # n, interface, group, signal, dir, pull, drive, function-index, remark
 (10, "GPIO", "USB20_0", "CPU_USB20_0_CTRL_1V8:USB0_OC_B",  "0: IN",  "3: PULL_UP",   "1: DRV_4_MA", 0, None),
 (11, "GPIO", "USB20_0", "CPU_USB20_0_CTRL_1V8:USB0_PWR_EN", "1: OUT", "0: NO_PULL",  "1: DRV_4_MA", 0, "PD during reset"),
 (12, "GPIO", "USB20_1", "CPU_USB20_1_CTRL_1V8:USB1_OC_B",  "0: IN",  "3: PULL_UP",   "1: DRV_4_MA", 0, "External pull-down required on this line"),   # deliberate contradiction for the AI data check
 (13, "GPIO", "USB20_1", "CPU_USB20_1_CTRL_1V8:USB1_PWR_EN", "1: OUT", "0: NO_PULL",  "1: DRV_4_MA", 0, None),
 (14, "GPIO", "USBHUB",  "CPU_USBHUB_CTRL_1V8:HUB_RST",     "1: OUT", "0: NO_PULL",  "3: DRV_8_MA", 0, None),
 (15, "GPIO", "BOOT",    "SA6150_BOOTCFG_1V8:FORCE_USB",    "0: IN",  "1: PULL_DOWN", "0: DRV_2_MA", 0, "PD: FORCED_USB pin"),
 (20, "HSUART_QUP13", "Bluetooth", "CPU_BT_UART_TX", "1: OUT", "0: NO_PULL", "0: DRV_2_MA", 1, None),
 (21, "HSUART_QUP13", "Bluetooth", "CPU_BT_UART_RX", "0: IN",  "3: PULL_UP",  "0: DRV_2_MA", 1, None),
 (22, "I2C_QUP2", "Camera", "CAM_I2C_SDA", "1: OUT", "0: NO_PULL", "1: DRV_4_MA", 2, None),
 (23, "I2C_QUP2", "Camera", "CAM_I2C_SCL", "1: OUT", "0: NO_PULL", "1: DRV_4_MA", 2, None),
]
for p in PINS:
    row(R, *p); R += 1
out = os.path.join(HERE, "demo_workbook.xlsx"); wb.save(out); print("wrote", out, f"({len(PINS)} pins)")
