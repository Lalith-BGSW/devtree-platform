"""
Parses the PinMux workbook (.xlsm/.xlsx).

Two-phase design:
  1. parse_workbook() extracts row-2/row-3 headers PURELY structurally -
     entity/base_field are split from the "Field (Entity)" label pattern
     itself (a generic regex, not a schema lookup), so this phase works
     on ANY workbook shaped like this, known or not.
  2. build_analysis() assigns FINAL roles: tries an exact match against
     every schema in the registry (app/schema_registry.py) first - if
     matched, roles/categories come from that confirmed schema, zero
     guessing, same as the original hardcoded-only design. If nothing
     matches, entities are classified using structural pattern detection
     (sub-groups: fields shaped like Port*/Instance*/RegSetting*, works
     for any count/naming) plus keyword-based SUGGESTIONS for SoC vs
     project - explicitly marked unconfirmed, never used for real parsing
     until a human confirms via /classify-schema, which saves it as a new
     registry entry for next time.

Still reads BOTH visible and hidden columns (hidden columns are never
skipped) and preserves visibility status per column, per the original spec.
"""

from __future__ import annotations
import re
import uuid
import openpyxl
from openpyxl.utils import get_column_letter

from app.models import ColumnMeta, EntitySummary, WorkbookAnalysis
from app.schema_registry import (
    match_known_schema, detect_subgroup_shaped_entities, suggest_role,
)

_LABEL_PATTERN = re.compile(r"^(.*?)\s*\(([^()]+)\)\s*$")


class WorkbookParseError(Exception):
    pass


class ParsedWorkbook:
    """In-memory representation of one parsed workbook, keyed by workbook_id."""

    def __init__(self, workbook_id: str, sheet_name: str, header_row: int,
                 data_start_row: int, columns: list[ColumnMeta],
                 raw_rows: list[list]):
        self.workbook_id = workbook_id
        self.sheet_name = sheet_name
        self.header_row = header_row
        self.data_start_row = data_start_row
        self.columns = columns
        self.raw_rows = raw_rows  # list of rows, each a list aligned to `columns`


def _split_label(field_label: str | None) -> tuple[str | None, str | None]:
    """Splits "Base Field (Entity)" into (base_field, entity) via a
    generic regex - purely structural, works regardless of what schema
    (if any) the entity/field names match. Returns (field_label, None)
    if the label doesn't have the "(...)" suffix shape at all."""
    if not field_label:
        return field_label, None
    m = _LABEL_PATTERN.match(field_label)
    if m:
        return m.group(1).strip(), m.group(2).strip()
    return field_label, None


def parse_workbook(file_path: str, sheet_name: str = "Database",
                    group_header_row: int = 2, field_header_row: int = 3
                    ) -> ParsedWorkbook:
    """
    Reads the workbook and returns a ParsedWorkbook with every column's
    STRUCTURAL metadata (visible/hidden, raw entity/base_field split from
    the label text itself) and the raw data rows. Final role/category
    assignment happens in build_analysis(), not here - see module
    docstring. Uses keep_vba=True so macro-enabled .xlsm files load
    without stripping the VBA project.
    """
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True, keep_vba=True)
    except Exception as e:
        raise WorkbookParseError(f"Could not open workbook: {e}") from e

    if sheet_name not in wb.sheetnames:
        # Fall back to the first sheet rather than hard-failing, but this
        # will show up as a schema mismatch in build_analysis() since none
        # of the expected columns will be found.
        sheet_name = wb.sheetnames[0]

    ws = wb[sheet_name]
    max_col = ws.max_column
    max_row = ws.max_row

    columns: list[ColumnMeta] = []
    for col_idx in range(1, max_col + 1):
        letter = get_column_letter(col_idx)
        group_val = ws.cell(row=group_header_row, column=col_idx).value
        field_val = ws.cell(row=field_header_row, column=col_idx).value
        group_label = str(group_val).strip() if group_val is not None else None
        field_label = str(field_val).strip() if field_val is not None else None

        dim = ws.column_dimensions.get(letter)
        is_hidden = bool(dim.hidden) if dim is not None else False

        raw_base_field, raw_entity = _split_label(field_label)

        columns.append(ColumnMeta(
            index=col_idx - 1,
            excel_letter=letter,
            group_label=group_label,
            field_label=field_label,
            base_field=raw_base_field,
            entity=raw_entity,
            slug=None,       # filled in by build_analysis() once role is resolved
            category=None,   # filled in by build_analysis() once role is resolved
            is_hidden=is_hidden,
            known=False,     # filled in by build_analysis() once role is resolved
        ))

    data_start_row = field_header_row + 1
    raw_rows: list[list] = []
    for row_idx in range(data_start_row, max_row + 1):
        row_vals = [ws.cell(row=row_idx, column=c + 1).value for c in range(max_col)]
        if all(v is None or str(v).strip() == "" for v in row_vals):
            continue  # skip fully blank rows
        raw_rows.append(row_vals)

    return ParsedWorkbook(
        workbook_id=str(uuid.uuid4()),
        sheet_name=sheet_name,
        header_row=field_header_row,
        data_start_row=data_start_row,
        columns=columns,
        raw_rows=raw_rows,
    )


def build_analysis(pw: ParsedWorkbook) -> WorkbookAnalysis:
    """Tries to match this file's columns against every known schema in
    the registry (exact match only). If matched, assigns final entity/
    category/slug from that CONFIRMED schema - zero guessing, identical
    behavior to the original hardcoded-only design for a known file.

    If nothing matches, classifies entities using structural pattern
    detection (sub-groups) and keyword heuristics (SoC vs project) -
    explicitly marked as SUGGESTIONS (role_confirmed=False), and sets
    needs_classification=True. These suggestions are never used for real
    parsing (normalize(), generate_dts(), etc.) until a human confirms
    them via POST /classify-schema, which saves the confirmed structure
    as a new registry entry."""
    found_labels = {c.field_label for c in pw.columns if c.field_label}
    matched_schema = match_known_schema(found_labels)

    hidden_count = sum(1 for c in pw.columns if c.is_hidden)

    if matched_schema is not None:
        # Known file - assign final entity/category/slug/known from the
        # CONFIRMED schema, exactly as the original design did.
        for col in pw.columns:
            known_def = matched_schema.known_columns.get(col.field_label) if col.field_label else None
            if known_def:
                col.base_field = known_def.base_field
                col.entity = known_def.entity
                col.slug = known_def.slug
                col.category = known_def.category
                col.known = True

        by_entity: dict[str, list[ColumnMeta]] = {}
        for col in pw.columns:
            if col.entity:
                by_entity.setdefault(col.entity, []).append(col)

        def _role(entity: str) -> str:
            if entity == matched_schema.soc_entity:
                return "soc"
            if entity in matched_schema.project_entities:
                return "project"
            if entity in matched_schema.subgroup_entities:
                return "subgroup"
            return "unknown"

        entities = [
            EntitySummary(
                name=entity, role=_role(entity),
                columns=[c.base_field for c in cols if c.base_field],
                column_count=len(cols), role_confirmed=True,
            )
            for entity, cols in by_entity.items()
        ]
        role_order = {"soc": 0, "project": 1, "subgroup": 2, "unknown": 3}
        entities.sort(key=lambda e: (role_order.get(e.role, 4), e.name))

        return WorkbookAnalysis(
            workbook_id=pw.workbook_id, sheet_name=pw.sheet_name,
            total_columns=len(pw.columns), hidden_columns=hidden_count,
            header_row=pw.header_row, data_start_row=pw.data_start_row,
            entities=entities, columns=pw.columns,
            schema_match=True, schema_issues=[],
            matched_schema_name=matched_schema.name, needs_classification=False,
        )

    # --- No known schema matched - build SUGGESTIONS, never applied yet ---
    by_entity: dict[str, list[ColumnMeta]] = {}
    for col in pw.columns:
        if col.entity:
            by_entity.setdefault(col.entity, []).append(col)

    entity_to_base_fields = {
        entity: [c.base_field for c in cols if c.base_field]
        for entity, cols in by_entity.items()
    }
    subgroup_entities = set(detect_subgroup_shaped_entities(entity_to_base_fields))

    entities: list[EntitySummary] = []
    for entity, cols in by_entity.items():
        base_fields = [c.base_field for c in cols if c.base_field]
        if entity in subgroup_entities:
            role, reason = "subgroup", "matches Port/Instance/RegSetting register-slot shape"
        else:
            role, reason = suggest_role(base_fields)
        entities.append(EntitySummary(
            name=entity, role=role, columns=base_fields, column_count=len(cols),
            role_confirmed=False, suggestion_reason=reason,
        ))
    role_order = {"soc": 0, "project": 1, "subgroup": 2, "unknown": 3}
    entities.sort(key=lambda e: (role_order.get(e.role, 4), e.name))

    issues = [
        "This file's columns don't exactly match any known, previously-confirmed "
        "schema. The entity roles below are SUGGESTIONS based on field-name "
        "patterns, not confirmed - review and confirm them via the classification "
        "step before this workbook can be built.",
    ]

    return WorkbookAnalysis(
        workbook_id=pw.workbook_id, sheet_name=pw.sheet_name,
        total_columns=len(pw.columns), hidden_columns=hidden_count,
        header_row=pw.header_row, data_start_row=pw.data_start_row,
        entities=entities, columns=pw.columns,
        schema_match=False, schema_issues=issues,
        matched_schema_name=None, needs_classification=True,
    )
