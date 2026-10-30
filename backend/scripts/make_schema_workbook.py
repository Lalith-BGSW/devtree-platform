"""
Builds a synthetic .xlsx using the EXACT column labels from the hardcoded
schema (app/schema.py) - which were themselves verified byte-for-byte
against the real Nissan_CCS20_SA8155PSIP_PinMux.xlsx headers. Guarantees
a perfect schema match by construction, since it's generated FROM the
schema. All pin data below is fabricated - no real project data is used
anywhere in this file.

Includes deliberately awkward cases to stress-test the pipeline:
  - A UART TX/RX pair (should merge into one DTS node)
  - An I2C pin
  - DNC/AGND/NC "non-signal" rows (should be excluded from DTS + duplicate
    checks, per app/constants.py)
  - A genuine duplicate ball_id (should still be caught as a real error)
  - An embedded legend row (_PARA_ROW="DRV_STRENGTH") that should be
    filtered out entirely, never becoming a record
  - An "_END_ROW" marker row that IS still real data and should NOT be
    filtered
"""
import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import openpyxl
from openpyxl.utils import get_column_letter
from app.schema import KNOWN_COLUMNS

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Database"

# Build header rows 2 (group) and 3 (field) directly from the schema, in
# schema dict order, and mark the columns the real file has hidden.
REAL_HIDDEN_LABELS = {
    "SA8155SIP Ball Name (SA8155P)": False, "SA8155SIP Ball ID (SA8155P)": False,
    "InstanceEGPIO (PinMuxEGPIO)": True, "Instance1 (PinMux01)": True,
    "InstanceQUP1 (PinMuxQ1)": True, "RegSettingQUP1 (PinMuxQ1)": True,
    "SW driver names (CCS20)": True, "ResetValue (CCS20)": True,
    "External Termination (CCS20)": True, "IO_GPIO_PULL (CCS20)": True,
    "IO_DRV_STRENGTH (CCS20)": True,
}

labels = list(KNOWN_COLUMNS.keys())
label_to_col: dict[str, int] = {}
for i, label in enumerate(labels, start=1):
    col_def = KNOWN_COLUMNS[label]
    group = col_def.entity if col_def.entity else ("Parameter Class" if label == "_PARA_ROW" else "_END_SB")
    ws.cell(row=2, column=i, value=group)
    ws.cell(row=3, column=i, value=label)
    if REAL_HIDDEN_LABELS.get(label):
        ws.column_dimensions[get_column_letter(i)].hidden = True
    label_to_col[label] = i


def set_row(row_idx: int, values: dict[str, object]) -> None:
    for label, val in values.items():
        col = label_to_col[label]
        ws.cell(row=row_idx, column=col, value=val)


ROW = 4

# --- Row: UART TX ---
set_row(ROW, {
    "SA8155SIP Ball Name (SA8155P)": "GPIO_43", "SA8155SIP Ball ID (SA8155P)": "AF12",
    "Power Domain (SA8155P)": "VDD_1P8",
    "QUPPort1 (PinMuxQ1)": "QUP13", "InstanceQUP1 (PinMuxQ1)": "13", "RegSettingQUP1 (PinMuxQ1)": "0x21",
    "PinMuxed Name (CCS20)": "GPIO43", "Selected Interface (CCS20)": "HSUART_QUP13",
    "PinMux Setting (CCS20)": "PinMux01", "Gen5 Signal Name (CCS20)": "CPU_BT_UART_TX",
    "Interface Group (CCS20)": "Bluetooth", "Direction (CCS20)": "out",
    "IO_DRV_STRENGTH (CCS20)": "0: DRV_2_MA", "IO_GPIO_PULL (CCS20)": "3: PULL_UP",
})
ROW += 1

# --- Row: UART RX ---
set_row(ROW, {
    "SA8155SIP Ball Name (SA8155P)": "GPIO_44", "SA8155SIP Ball ID (SA8155P)": "AF13",
    "Power Domain (SA8155P)": "VDD_1P8",
    "QUPPort1 (PinMuxQ1)": "QUP13", "InstanceQUP1 (PinMuxQ1)": "13", "RegSettingQUP1 (PinMuxQ1)": "0x22",
    "PinMuxed Name (CCS20)": "GPIO44", "Selected Interface (CCS20)": "HSUART_QUP13",
    "PinMux Setting (CCS20)": "PinMux01", "Gen5 Signal Name (CCS20)": "CPU_BT_UART_RX",
    "Interface Group (CCS20)": "Bluetooth", "Direction (CCS20)": "in",
    "IO_DRV_STRENGTH (CCS20)": "0: DRV_2_MA", "IO_GPIO_PULL (CCS20)": "0: NO_PULL",
})
ROW += 1

# --- Row: I2C ---
set_row(ROW, {
    "SA8155SIP Ball Name (SA8155P)": "GPIO_21", "SA8155SIP Ball ID (SA8155P)": "AC09",
    "Power Domain (SA8155P)": "VDD_1P8",
    "PinMuxed Name (CCS20)": "GPIO21", "Selected Interface (CCS20)": "I2C_QUP2",
    "PinMux Setting (CCS20)": "PinMux02", "Gen5 Signal Name (CCS20)": "CAM_I2C_SDA",
    "Interface Group (CCS20)": "Camera", "Direction (CCS20)": "out",
    "IO_DRV_STRENGTH (CCS20)": "1: DRV_4_MA", "IO_GPIO_PULL (CCS20)": "0: NO_PULL",
})
ROW += 1

# --- 3 AGND rows sharing the same PinMuxed Name label (real workbook pattern) ---
for i in range(3):
    set_row(ROW, {
        "SA8155SIP Ball Name (SA8155P)": "AGND", "SA8155SIP Ball ID (SA8155P)": f"GND{i}",
        "PinMuxed Name (CCS20)": "AGND", "Selected Interface (CCS20)": "Power",
    })
    ROW += 1

# --- 2 DNC rows ---
for i in range(2):
    set_row(ROW, {
        "SA8155SIP Ball Name (SA8155P)": "NC", "SA8155SIP Ball ID (SA8155P)": f"NC{i}",
        "PinMuxed Name (CCS20)": "DNC", "Selected Interface (CCS20)": "DNC",
    })
    ROW += 1

# --- Genuine duplicate: same ball_id "AF12" reused by mistake ---
set_row(ROW, {
    "SA8155SIP Ball Name (SA8155P)": "GPIO_43", "SA8155SIP Ball ID (SA8155P)": "AF12",
    "PinMuxed Name (CCS20)": "GPIO43_DUP", "Selected Interface (CCS20)": "I2C_QUP5",
    "Gen5 Signal Name (CCS20)": "SOME_OTHER_SIGNAL",
})
ROW += 1

# --- Embedded legend row - should be filtered out entirely ---
set_row(ROW, {"_PARA_ROW": "DRV_STRENGTH", "ID (SA8155P)": "0: DRV_2_MA"})
ROW += 1

# --- _END_ROW marker - IS real data, must NOT be filtered ---
set_row(ROW, {
    "_PARA_ROW": "_END_ROW", "ID (SA8155P)": "999",
    "SA8155SIP Ball Name (SA8155P)": "GND_LAST", "SA8155SIP Ball ID (SA8155P)": "ZZ99",
    "PinMuxed Name (CCS20)": "GND", "Selected Interface (CCS20)": "Power",
})
ROW += 1

out_path = os.path.join(os.path.dirname(__file__), "schema_synthetic.xlsx")
wb.save(out_path)
print(f"wrote {out_path} ({ROW - 4} data rows, {len(labels)} columns from schema)")
